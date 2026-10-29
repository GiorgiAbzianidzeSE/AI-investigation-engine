import io

import pytest

from app.ingest.chunking import chunk_fixed, chunk_structured
from app.ingest.normalize import content_hash, normalize_text, slugify
from app.ingest.parsers import UnsupportedFormat, extract_date, guess_doc_type, parse_file


# ---------- normalize ----------
def test_normalize_collapses_whitespace_and_keeps_paragraphs():
    raw = "Hello\t  world  \r\n\r\n\r\n\r\nSecond\u00a0para  "
    assert normalize_text(raw) == "Hello world\n\nSecond para"


def test_hash_ignores_case_and_whitespace():
    assert content_hash("Hello   World\n\nfoo") == content_hash("hello world foo")
    assert content_hash("hello world") != content_hash("hello world!")


def test_slugify():
    assert slugify("INC-2041-incident-report.pdf") == "inc-2041-incident-report-pdf"


# ---------- parsers ----------
def test_markdown_parser():
    doc = parse_file("fraud-scorer-runbook.md", b"# Fraud runbook\n\nUpdated 2026-03-05.\n\n## Rollback\n\nSteps.")
    assert doc.title == "Fraud runbook"
    assert doc.format == "md"
    assert doc.date == "2026-03-05"
    assert doc.doc_type == "runbook"
    assert "## Rollback" in doc.text


def test_html_parser_headings_and_noise():
    html = b"""<html><head><title>Status page</title><style>x{}</style></head><body>
    <nav>menu</nav><h2>INC-2041</h2><p>Resolved on March 5, 2026.</p><ul><li>item</li></ul>
    <script>alert(1)</script><footer>foot</footer></body></html>"""
    doc = parse_file("status-page-march-2026.html", html)
    assert doc.title == "Status page"
    assert "# INC-2041" in doc.text
    assert "- item" in doc.text
    assert "menu" not in doc.text and "alert" not in doc.text and "foot" not in doc.text
    assert doc.date == "2026-03-05"
    assert doc.doc_type == "status_page"


def test_docx_parser_headings():
    import docx

    d = docx.Document()
    d.core_properties.title = "CHG-0311 change request"
    d.add_heading("Description", level=1)
    d.add_paragraph("Reduce the pool on 2026-03-03.")
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_file("CHG-0311-change-request.docx", buf.getvalue())
    assert doc.title == "CHG-0311 change request"
    assert "# Description" in doc.text
    assert doc.doc_type == "change_request"
    assert doc.date == "2026-03-03"


def test_pdf_parser():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_title("INC-2041 incident report")
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 6, "Card authorizations failed on 2026-03-05. " + "Payments were slow. " * 20, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 13)
    pdf.multi_cell(0, 8, "Root cause", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 6, "The pool was exhausted.", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.multi_cell(0, 6, "Second paragraph.", new_x="LMARGIN", new_y="NEXT")
    doc = parse_file("INC-2041-incident-report.pdf", bytes(pdf.output()))
    assert doc.title == "INC-2041 incident report"
    assert doc.date == "2026-03-05"
    assert doc.doc_type == "incident"
    assert "\n# Root cause\n" in doc.text
    assert "The pool was exhausted.\n\nSecond paragraph." in doc.text
    assert chunk_structured(doc.text, doc.title)[-1].heading == "Root cause"


def test_unsupported_and_empty():
    with pytest.raises(UnsupportedFormat):
        parse_file("x.exe", b"abc")
    with pytest.raises(ValueError):
        parse_file("x.txt", b"   \n  ")


def test_date_and_doc_type_helpers():
    assert extract_date("On March 19, 2026 payouts") == "2026-03-19"
    assert extract_date("no date") is None
    assert guess_doc_type("PM-2041-postmortem.pdf") == "postmortem"
    assert guess_doc_type("random.txt") == "document"


# ---------- chunking ----------
def test_fixed_window_sizes_and_overlap():
    text = " ".join(f"w{i}" for i in range(450))
    chunks = chunk_fixed(text, size=200, overlap=40)
    assert [len(c.text.split()) for c in chunks] == [200, 200, 130]
    assert chunks[1].text.split()[0] == "w160"
    assert chunks[-1].text.split()[-1] == "w449"


def test_fixed_rejects_bad_overlap():
    with pytest.raises(ValueError):
        chunk_fixed("a b c", size=10, overlap=10)


def test_structured_prefix_and_sections():
    text = "Intro para.\n\n# Impact\n\nCustomers saw failures.\n\n# Timeline\n\n09:40 alert."
    chunks = chunk_structured(text, "INC-2041", size=200)
    assert [c.heading for c in chunks] == [None, "Impact", "Timeline"]
    assert chunks[1].text == "INC-2041 > Impact: Customers saw failures."
    assert chunks[0].text.startswith("INC-2041: ")


def test_structured_never_cuts_sentences():
    sentences = [f"Sentence number {i} has exactly seven words." for i in range(60)]
    text = "# Long\n\n" + " ".join(sentences)  # one 420-word paragraph -> split by sentence
    chunks = chunk_structured(text, "T", size=50)
    assert len(chunks) > 1
    for c in chunks:
        body = c.text.split(": ", 1)[1]
        assert body.endswith("words.")
        assert len(body.split()) <= 50
    assert sum(c.text.count("Sentence number") for c in chunks) == 60
